# SPDX-FileCopyrightText: Copyright (c) 2024 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

import os
os.environ.setdefault("USER_AGENT", "project-agentic-rag/1.0 (+https://yourdomain)")

from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import (
    WebBaseLoader,
    UnstructuredPDFLoader,
    TextLoader,
    CSVLoader,
    UnstructuredWordDocumentLoader,
    UnstructuredPowerPointLoader,
    UnstructuredExcelLoader,
    UnstructuredHTMLLoader,
    UnstructuredEmailLoader,
)
from langchain_community.vectorstores import Chroma
import json
import urllib.request

from typing import Any, Dict, List, Tuple, Union
from urllib.parse import urlparse
import shutil
import mimetypes
import time



class StructuredXLSXLoader:
    """
    XLSX loader that preserves sheet/row/column structure as markdown tables.

    Each worksheet becomes a section with a "## Sheet: <name>" header followed
    by a markdown table. Column letters are preserved in a companion line so
    queries like "what is in cell B2?" remain answerable.
    """

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self):
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        from langchain_core.documents import Document

        wb = load_workbook(self.file_path, data_only=True, read_only=True)
        parts = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append([("" if c is None else str(c)) for c in row])

            if not rows:
                continue

            parts.append(f"## Sheet: {sheet_name}")

            # Detect widest row to define column count
            ncols = max(len(r) for r in rows) if rows else 0
            if ncols == 0:
                continue

            # Column letter reference line — makes "cell B2" queries resolvable
            col_letters = [get_column_letter(i + 1) for i in range(ncols)]
            parts.append("Columns: " + " | ".join(col_letters))

            # First row as header, rest as data
            header = rows[0] + [""] * (ncols - len(rows[0]))
            parts.append("| " + " | ".join(header) + " |")
            parts.append("|" + "|".join(["---"] * ncols) + "|")

            for r in rows[1:]:
                r = r + [""] * (ncols - len(r))
                parts.append("| " + " | ".join(r) + " |")

            parts.append("")  # blank line between sheets

        body = "\n".join(parts).strip() if parts else "(empty workbook)"
        return [Document(
            page_content=body,
            metadata={"source": self.file_path, "format": "xlsx"},
        )]


class StructuredDocxLoader:
    """
    DOCX loader that preserves heading hierarchy, paragraphs, lists, and
    tables as markdown. Headings become "# text" markers, lists become
    "- item" or "1. item", tables become pipe tables. Structure survives
    so the LLM can reason about sections and hierarchy.
    """

    def __init__(self, file_path: str):
        self.file_path = file_path

    def _heading_level(self, style_name: str):
        if not style_name:
            return None
        s = style_name.lower()
        if s.startswith("heading"):
            parts = s.replace("heading", "").strip()
            try:
                return int(parts)
            except ValueError:
                return 1
        return None

    def load(self):
        from docx import Document as DocxDocument
        from langchain_core.documents import Document

        doc = DocxDocument(self.file_path)
        parts = []

        for element in doc.element.body:
            tag = element.tag.split("}")[-1]

            if tag == "p":
                para = next((p for p in doc.paragraphs if p._element is element), None)
                if para is None:
                    continue
                text = para.text.strip()
                if not text:
                    continue
                level = self._heading_level(para.style.name if para.style else "")
                if level is not None:
                    parts.append("#" * level + " " + text)
                elif para.style and "List Bullet" in para.style.name:
                    parts.append("- " + text)
                elif para.style and "List Number" in para.style.name:
                    parts.append("1. " + text)
                else:
                    parts.append(text)

            elif tag == "tbl":
                tbl = next((t for t in doc.tables if t._element is element), None)
                if tbl is None:
                    continue
                rows = []
                for row in tbl.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    rows.append(cells)
                if not rows:
                    continue
                ncols = max(len(r) for r in rows)
                header = rows[0] + [""] * (ncols - len(rows[0]))
                parts.append("| " + " | ".join(header) + " |")
                parts.append("|" + "|".join(["---"] * ncols) + "|")
                for r in rows[1:]:
                    r = r + [""] * (ncols - len(r))
                    parts.append("| " + " | ".join(r) + " |")
                parts.append("")

        body = "\n\n".join(parts) if parts else "(empty document)"
        return [Document(
            page_content=body,
            metadata={"source": self.file_path, "format": "docx"},
        )]


class StructuredHTMLLoader:
    """
    HTML loader that preserves document structure as markdown.

    Extracts headings (h1-h6), paragraphs, list items, and tables, emitting
    them in document order with markdown markers (# for headings, - for list
    items, | for tables). This keeps structural queries answerable — for
    example "What text is in the h1?" — because the tag identity is preserved
    as a markdown marker that both the retriever and the LLM can see.
    """

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self):
        from bs4 import BeautifulSoup
        from langchain_core.documents import Document

        with open(self.file_path, "r", encoding="utf-8", errors="ignore") as f:
            html = f.read()

        soup = BeautifulSoup(html, "html.parser")
        for t in soup(["script", "style", "noscript"]):
            t.decompose()

        parts = []
        for element in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6",
                                      "p", "li", "table"]):
            # Skip inline elements inside tables — handled at table level
            if element.name != "table" and element.find_parent("table"):
                continue

            if element.name.startswith("h"):
                level = int(element.name[1])
                text = element.get_text(strip=True)
                if text:
                    parts.append("#" * level + " " + text)

            elif element.name == "p":
                text = element.get_text(strip=True)
                if text:
                    parts.append(text)

            elif element.name == "li":
                text = element.get_text(strip=True)
                if text:
                    parts.append("- " + text)

            elif element.name == "table":
                rows = []
                for tr in element.find_all("tr"):
                    cells = [c.get_text(strip=True)
                             for c in tr.find_all(["td", "th"])]
                    if cells:
                        rows.append(cells)
                if rows:
                    header = rows[0]
                    parts.append("| " + " | ".join(header) + " |")
                    parts.append("|" + "|".join(["---"] * len(header)) + "|")
                    for row in rows[1:]:
                        while len(row) < len(header):
                            row.append("")
                        parts.append("| " + " | ".join(row) + " |")

        body = "\n\n".join(parts) if parts else "(empty document)"

        return [Document(
            page_content=body,
            metadata={"source": self.file_path, "format": "html"},
        )]


# Local Ollama embedding configuration
EMBEDDINGS_MODEL = os.getenv("OLLAMA_EMBED_MODEL", "qwen3-embedding:0.6b")
OLLAMA_EMBED_URL = os.getenv(
    "OLLAMA_EMBED_URL",
    "http://host.docker.internal:11434/api/embed",
)

print(f"[config] Using local Ollama embedding model: {EMBEDDINGS_MODEL}")
print(f"[config] Ollama embedding endpoint: {OLLAMA_EMBED_URL}")


class OllamaEmbeddings:
    """Minimal embedding adapter for Ollama."""

    def __init__(self, model=EMBEDDINGS_MODEL, url=OLLAMA_EMBED_URL):
        self.model = model
        self.url = url

    def _embed(self, inputs):
        # Ollama supports batched embedding requests. Keep batches bounded to
        # avoid oversized/bulk-request failures during large RAG ingestions.
        batch_size = 32
        all_embeddings = []

        for start in range(0, len(inputs), batch_size):
            batch = inputs[start:start + batch_size]

            payload = {
                "model": self.model,
                "input": batch,
            }

            request = urllib.request.Request(
                self.url,
                data=json.dumps(payload).encode("utf-8"),
                headers={"Content-Type": "application/json"},
                method="POST",
            )

            try:
                with urllib.request.urlopen(request, timeout=120) as response:
                    result = json.loads(response.read().decode("utf-8"))
            except Exception as e:
                raise RuntimeError(
                    f"Ollama embedding request failed for batch "
                    f"{start + 1}-{start + len(batch)}: {e}"
                ) from e

            embeddings = result.get("embeddings")
            if not embeddings:
                raise RuntimeError(
                    f"Ollama embedding response missing embeddings "
                    f"for batch {start + 1}-{start + len(batch)}: {result}"
                )

            all_embeddings.extend(embeddings)

        return all_embeddings

    def embed_documents(self, texts):
        return self._embed(list(texts))

    def embed_query(self, text):
        return self._embed([text])[0]


# Set the chunk size and overlap for the text splitter.
DEFAULT_CHUNK_SIZE = 250
DEFAULT_CHUNK_OVERLAP = 0

CHUNK_SIZE = int(os.getenv("CHUNK_SIZE", DEFAULT_CHUNK_SIZE))
CHUNK_OVERLAP = int(os.getenv("CHUNK_OVERLAP", DEFAULT_CHUNK_OVERLAP))


# Adding nltk data
import nltk

def download_nltk_if_missing():
    try:
        nltk.data.find('tokenizers/punkt')
    except LookupError:
        nltk.download('punkt')

    try:
        nltk.data.find('taggers/averaged_perceptron_tagger')
    except LookupError:
        nltk.download('averaged_perceptron_tagger')

download_nltk_if_missing()

    
# nltk.download("punkt")
# nltk.download("averaged_perceptron_tagger")

# Functions for dealing with URLs
def is_valid_url(url: str) -> bool:
    """ This is a helper function for checking if the URL is valid. It isn't fail proof, but it will catch most common errors. """
    try:
        result = urlparse(url)
        return all([result.scheme, result.netloc])
    except:
        return False

def safe_load(url):
    """ This is a helper function for loading the URL. It protects against false negatives from is_value_url and 
        filters for actual web pages. Returns None if it fails.
    """
    try:
        return WebBaseLoader(url).load()
    except Exception as e:
        print(f"[Documents] ⚠ Could not load {url}: {e}")
        return None


def upload(urls: List[str]):
    """ This is a helper function for parsing the user inputted URLs and uploading them into the vector store. """

    urls = [url.strip() for url in urls if url.strip()]
    for url in urls:
        if not is_valid_url(url):
            print(f"[Documents] ⚠ Skipping invalid URL: {url}")
    urls = [url for url in urls if is_valid_url(url)]

    if not urls:
        print("[Documents] ✗ No valid URLs provided — nothing was added to the context.")
        return None

    print(f"[Documents] Adding {len(urls)} webpage(s) to the context...")
    docs = []
    for url in urls:
        result = safe_load(url)
        if result is not None:
            docs.append(result)
            print(f"[Documents] ✓ Loaded {url}")

    docs_list = [item for sublist in docs for item in sublist]

    if not docs_list:
        # If no documents were loaded, return None
        print("[Documents] ✗ None of the webpages could be loaded — nothing was added to the context.")
        return None

    print(f"[Documents] Loaded {len(docs)} of {len(urls)} webpage(s) successfully")

    try:
        doc_splits = split_documents(docs_list)
        return embed_documents(doc_splits)

    except Exception as e:
        print(f"[Documents] ✗ Failed to add webpages to the context: {e}")
        return None


# Functions for dealing with file uploads/embeddings

## Helper functions

def load_documents_from_files(file_paths: List[str]) -> List[Any]:
    """Load and return documents from supported file types."""
    docs = []

    for fpath in file_paths:
        ext = os.path.splitext(fpath)[-1].lower()

        loader_cls = {
            ".pdf": UnstructuredPDFLoader,
            ".txt": TextLoader,
            ".md": TextLoader,
            ".csv": CSVLoader,
            ".docx": StructuredDocxLoader,
            # .doc (Word 97-2003) intentionally omitted — python-docx only reads .docx.
            # Legacy .doc requires LibreOffice which is not installed in this container.
            ".pptx": UnstructuredPowerPointLoader,
            ".xlsx": StructuredXLSXLoader,
            ".html": StructuredHTMLLoader,
            ".htm": StructuredHTMLLoader,
            ".eml": UnstructuredEmailLoader,
        }.get(ext)

        if loader_cls is None:
            print(f"[Documents] ⚠ Skipping unsupported file type: {os.path.basename(fpath)} (supported: .pdf, .txt, .md, .csv, .docx, .pptx, .xlsx, .html, .htm, .eml)")
            continue

        try:
            loaded = loader_cls(fpath).load()
            docs.append(loaded)
            print(f"[Documents] ✓ Loaded {os.path.basename(fpath)}")
        except Exception as e:
            print(f"[Documents] ✗ Failed to load {os.path.basename(fpath)}: {e}")

    return [item for sublist in docs for item in sublist]


def split_documents(docs: List[Any]):
    """Split documents into smaller chunks using recursive splitter."""
    print(f"[Documents] Splitting {len(docs)} document(s) into chunks of ~{CHUNK_SIZE} tokens ({CHUNK_OVERLAP} token overlap)...")

    splitter = RecursiveCharacterTextSplitter.from_tiktoken_encoder(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP
    )
    doc_splits = splitter.split_documents(docs)
    print(f"[Documents] Created {len(doc_splits)} chunks")
    return doc_splits

def embed_documents(doc_splits: List[Any]):
    """Embed and store the split documents into Chroma vectorstore."""
    try:
        print(f"[Documents] Embedding {len(doc_splits)} chunks with {EMBEDDINGS_MODEL}...")
        start = time.time()

        vectorstore = Chroma.from_documents(
            documents=doc_splits,
            collection_name="rag-chroma",
            embedding=OllamaEmbeddings(model=EMBEDDINGS_MODEL),
            persist_directory="/project/data",
        )
        print(f"[Documents] ✓ Context ready — {len(doc_splits)} chunks stored in the vector database ({time.time() - start:.1f}s)")
        return vectorstore

    except Exception as e:
        print(f"[Documents] ✗ Embedding failed — could not build the context: {e}")
        return None


## Main function that use helper functions 

def upload_files(file_paths: List[str]):
    """Upload files into the vector store pipeline."""

    if not file_paths:
        print("[Documents] ✗ No files provided — nothing was added to the context.")
        return None

    print(f"[Documents] Adding {len(file_paths)} file(s) to the context...")

    try:
        docs_list = load_documents_from_files(file_paths)

        if not docs_list:
            print("[Documents] ✗ None of the files could be loaded — nothing was added to the context.")
            return None

        doc_splits = split_documents(docs_list)
        return embed_documents(doc_splits)

    except Exception as e:
        print(f"[Documents] ✗ Failed to add files to the context: {e}")
        return None




def _clear(
    persist_directory: str = "/project/data",
    collection_name: str = "rag-chroma",
    delete_all: bool = True
):
    """Clear the Chroma collection and optionally delete all shard folders (excluding hidden files)."""
    try:
        print("[Documents] Clearing the context...")

        # Clear the collection via Chroma client
        vectorstore = Chroma(
            collection_name=collection_name,
            embedding_function=OllamaEmbeddings(model=EMBEDDINGS_MODEL),
            persist_directory=persist_directory,
        )
        vectorstore._client.delete_collection(name=collection_name)
        vectorstore._client.create_collection(name=collection_name)

        if delete_all:
            removed = 0
            for item in os.listdir(persist_directory):
                if item.startswith(".") or item.startswith("chroma.") or item.startswith("readme-images"):
                    continue  # Skip hidden files like .gitkeep and the sqlite3 file

                path = os.path.join(persist_directory, item)
                try:
                    if os.path.isfile(path):
                        os.remove(path)
                        removed += 1
                    elif os.path.isdir(path):
                        shutil.rmtree(path)
                        removed += 1
                except Exception as file_err:
                    print(f"[Documents] ⚠ Could not delete cached index data '{item}': {file_err}")
            if removed:
                print(f"[Documents] Removed {removed} cached index item(s) from disk")

        print("[Documents] ✓ Context cleared — the vector database is now empty")

    except Exception as e:
        print(f"[Documents] ✗ Failed to clear the context: {e}")


      
# def clear():
#     """ This is a helper function for emptying the collection the vector store. """
#     vectorstore = Chroma(
#         collection_name="rag-chroma",
#         embedding_function=OllamaEmbeddings(model=EMBEDDINGS_MODEL),
#         persist_directory="/project/data",
#     )
    
#     vectorstore._client.delete_collection(name="rag-chroma")
#     vectorstore._client.create_collection(name="rag-chroma")

def get_retriever(): 
    """ This is a helper function for returning the retriever object of the vector store. """
    vectorstore = Chroma(
        collection_name="rag-chroma",
        embedding_function=OllamaEmbeddings(model=EMBEDDINGS_MODEL),
        persist_directory="/project/data",
    )
    retriever = vectorstore.as_retriever(search_kwargs={"k": 25})
    return retriever
